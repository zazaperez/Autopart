"""
MOTOPART — Exportación de reportes a Excel y PDF.

Requiere las librerías 'openpyxl' (Excel) y 'reportlab' (PDF):
    pip install openpyxl reportlab
"""

import io

# ============================================================
# Utilidades comunes
# ============================================================


def _fmt_cop(valor):
    """Formatea un número como pesos colombianos sin decimales: 1.234.567"""
    try:
        numero = int(round(float(valor)))
    except (TypeError, ValueError):
        return "0"
    negativo = numero < 0
    numero = abs(numero)
    texto = f"{numero:,}".replace(",", ".")
    return f"-{texto}" if negativo else texto


def _xlsx_response(wb, filename):
    from django.http import HttpResponse

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type=(
            "application/vnd.openxmlformats-officedocument" ".spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _pdf_response(buffer, filename):
    from django.http import HttpResponse

    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


COLOR_ROJO = "C0392B"
COLOR_GRIS_CLARO = "F8F9FB"


def _encabezado_excel(ws, fila, num_columnas):
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill(start_color=COLOR_ROJO, end_color=COLOR_ROJO, fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, num_columnas + 1):
        celda = ws.cell(row=fila, column=col)
        celda.fill = fill
        celda.font = font
        celda.alignment = Alignment(horizontal="center", vertical="center")


def _titulo_excel(ws, texto, fila=1):
    from openpyxl.styles import Font

    ws.cell(row=fila, column=1, value=texto).font = Font(bold=True, size=14)


def _ajustar_anchos(ws, anchos):
    from openpyxl.utils import get_column_letter

    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _estilo_titulo_pdf():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="MPTitulo",
            parent=styles["Title"],
            textColor=colors.HexColor("#1a1a2e"),
            fontSize=16,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MPSubtitulo",
            parent=styles["Normal"],
            textColor=colors.HexColor("#7a8290"),
            fontSize=10,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MPSeccion",
            parent=styles["Heading3"],
            textColor=colors.HexColor("#1a1a2e"),
            fontSize=12,
            spaceBefore=14,
        )
    )
    return styles


def _tabla_estilo_base():
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle

    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + COLOR_ROJO)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e0e0e0")),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.white, colors.HexColor("#fafbfc")],
            ),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
    )


# ============================================================
# 1. VENTAS POR PERIODO
# ============================================================


def exportar_ventas_excel(datos):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    _titulo_excel(ws, "MOTOPART — Reporte de Ventas por Periodo")
    ws["A2"] = f"Periodo: {datos['desde']} a {datos['hasta']}"

    resumen = [
        ("Total vendido", float(datos["total_ventas"])),
        ("Pedidos", datos["num_pedidos"]),
        ("Ticket promedio", float(datos["ticket_promedio"])),
        ("Subtotal sin IVA", float(datos["subtotal_sin_iva"])),
        ("IVA (19%)", float(datos["iva_total"])),
    ]
    fila = 4
    for etiqueta, valor in resumen:
        ws.cell(row=fila, column=1, value=etiqueta)
        ws.cell(row=fila, column=2, value=valor)
        fila += 1

    fila_tabla = fila + 1
    encabezados = ["Fecha", "Pedido", "Cliente", "Estado", "Total"]
    for i, h in enumerate(encabezados, start=1):
        ws.cell(row=fila_tabla, column=i, value=h)
    _encabezado_excel(ws, fila_tabla, len(encabezados))

    fila = fila_tabla + 1
    for p in datos["pedidos_detalle"]:
        ws.cell(row=fila, column=1, value=p.fecha_pedido.strftime("%d/%m/%Y %H:%M"))
        ws.cell(row=fila, column=2, value=f"#{p.pk}")
        ws.cell(
            row=fila, column=3, value=p.usuario.get_full_name() or p.usuario.username
        )
        ws.cell(row=fila, column=4, value=p.get_estado_display())
        ws.cell(row=fila, column=5, value=float(p.total))
        fila += 1

    if not datos["pedidos_detalle"]:
        ws.cell(row=fila, column=1, value="No hay ventas registradas en este periodo.")

    _ajustar_anchos(ws, [18, 10, 30, 16, 16])
    return _xlsx_response(wb, "reporte_ventas.xlsx")


def exportar_ventas_pdf(datos):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = _estilo_titulo_pdf()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, topMargin=1.5 * cm, bottomMargin=1.5 * cm
    )
    elementos = [
        Paragraph("MOTOPART — Reporte de Ventas por Periodo", styles["MPTitulo"]),
        Paragraph(
            f"Periodo: {datos['desde']} a {datos['hasta']}", styles["MPSubtitulo"]
        ),
        Spacer(1, 14),
    ]

    resumen_data = [
        ["Total vendido", f"${_fmt_cop(datos['total_ventas'])}"],
        ["Pedidos", str(datos["num_pedidos"])],
        ["Ticket promedio", f"${_fmt_cop(datos['ticket_promedio'])}"],
        ["Subtotal sin IVA", f"${_fmt_cop(datos['subtotal_sin_iva'])}"],
        ["IVA (19%) a declarar", f"${_fmt_cop(datos['iva_total'])}"],
    ]
    t_resumen = Table(resumen_data, colWidths=[7 * cm, 6 * cm])
    t_resumen.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#" + COLOR_GRIS_CLARO),
                ),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e0e0e0")),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elementos.append(t_resumen)
    elementos.append(Paragraph("Detalle de pedidos", styles["MPSeccion"]))

    tabla_datos = [["Fecha", "Pedido", "Cliente", "Estado", "Total"]]
    for p in datos["pedidos_detalle"]:
        tabla_datos.append(
            [
                p.fecha_pedido.strftime("%d/%m/%Y %H:%M"),
                f"#{p.pk}",
                p.usuario.get_full_name() or p.usuario.username,
                p.get_estado_display(),
                f"${_fmt_cop(p.total)}",
            ]
        )
    if len(tabla_datos) == 1:
        tabla_datos.append(["—", "—", "Sin ventas en este periodo", "—", "—"])

    t = Table(
        tabla_datos,
        colWidths=[3.2 * cm, 1.8 * cm, 5.5 * cm, 2.8 * cm, 2.7 * cm],
        repeatRows=1,
    )
    t.setStyle(_tabla_estilo_base())
    elementos.append(t)

    doc.build(elementos)
    return _pdf_response(buffer, "reporte_ventas.pdf")


# ============================================================
# 2. INVENTARIO / STOCK BAJO
# ============================================================


def exportar_inventario_excel(datos):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    _titulo_excel(ws, "MOTOPART — Reporte de Inventario / Stock bajo")
    sub = "Categoría: " + (datos["categoria_nombre"] or "Todas")
    if datos.get("solo_bajo"):
        sub += " — solo productos con stock bajo"
    ws["A2"] = sub

    resumen = [
        ("Productos activos", datos["total_productos"]),
        ("Con stock bajo", datos["total_stock_bajo"]),
        ("Sin stock (agotados)", datos["total_sin_stock"]),
        ("Valor total de inventario (sin IVA)", float(datos["valor_inventario_total"])),
    ]
    fila = 4
    for etiqueta, valor in resumen:
        ws.cell(row=fila, column=1, value=etiqueta)
        ws.cell(row=fila, column=2, value=valor)
        fila += 1

    fila_tabla = fila + 1
    encabezados = [
        "Producto",
        "SKU",
        "Categoría",
        "Stock actual",
        "Stock mínimo",
        "Estado",
        "Valor en stock (sin IVA)",
    ]
    for i, h in enumerate(encabezados, start=1):
        ws.cell(row=fila_tabla, column=i, value=h)
    _encabezado_excel(ws, fila_tabla, len(encabezados))

    fila = fila_tabla + 1
    for item in datos["productos"]:
        p = item["obj"]
        if p.stock == 0:
            estado = "Agotado"
        elif p.stock_bajo():
            estado = "Stock bajo"
        else:
            estado = "OK"
        ws.cell(row=fila, column=1, value=p.nombre)
        ws.cell(row=fila, column=2, value=p.sku)
        ws.cell(row=fila, column=3, value=p.categoria.nombre if p.categoria else "—")
        ws.cell(row=fila, column=4, value=p.stock)
        ws.cell(row=fila, column=5, value=p.stock_minimo)
        ws.cell(row=fila, column=6, value=estado)
        ws.cell(row=fila, column=7, value=float(item["valor_stock"]))
        if estado != "OK":
            for col in range(1, 8):
                ws.cell(row=fila, column=col).font = Font(color="C0392B")
        fila += 1

    if not datos["productos"]:
        ws.cell(
            row=fila, column=1, value="No hay productos con los filtros seleccionados."
        )

    _ajustar_anchos(ws, [30, 14, 18, 12, 12, 12, 20])
    return _xlsx_response(wb, "reporte_inventario.xlsx")


def exportar_inventario_pdf(datos):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = _estilo_titulo_pdf()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, topMargin=1.5 * cm, bottomMargin=1.5 * cm
    )
    sub = "Categoría: " + (datos["categoria_nombre"] or "Todas")
    if datos.get("solo_bajo"):
        sub += " — solo productos con stock bajo"

    elementos = [
        Paragraph("MOTOPART — Reporte de Inventario / Stock bajo", styles["MPTitulo"]),
        Paragraph(sub, styles["MPSubtitulo"]),
        Spacer(1, 14),
    ]

    resumen_data = [
        ["Productos activos", str(datos["total_productos"])],
        ["Con stock bajo", str(datos["total_stock_bajo"])],
        ["Sin stock (agotados)", str(datos["total_sin_stock"])],
        [
            "Valor total inventario (sin IVA)",
            f"${_fmt_cop(datos['valor_inventario_total'])}",
        ],
    ]
    t_resumen = Table(resumen_data, colWidths=[8 * cm, 6 * cm])
    t_resumen.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#" + COLOR_GRIS_CLARO),
                ),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e0e0e0")),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elementos.append(t_resumen)
    elementos.append(Paragraph("Detalle de productos", styles["MPSeccion"]))

    tabla_datos = [
        ["Producto", "SKU", "Categoría", "Stock", "Mínimo", "Estado", "Valor"]
    ]
    for item in datos["productos"]:
        p = item["obj"]
        if p.stock == 0:
            estado = "Agotado"
        elif p.stock_bajo():
            estado = "Stock bajo"
        else:
            estado = "OK"
        tabla_datos.append(
            [
                p.nombre,
                p.sku,
                p.categoria.nombre if p.categoria else "—",
                str(p.stock),
                str(p.stock_minimo),
                estado,
                f"${_fmt_cop(item['valor_stock'])}",
            ]
        )
    if len(tabla_datos) == 1:
        tabla_datos.append(
            ["—", "—", "Sin productos con estos filtros", "—", "—", "—", "—"]
        )

    t = Table(
        tabla_datos,
        colWidths=[
            4.2 * cm,
            2.2 * cm,
            2.6 * cm,
            1.6 * cm,
            1.6 * cm,
            2.2 * cm,
            2.4 * cm,
        ],
        repeatRows=1,
    )
    estilo = _tabla_estilo_base()
    for i, item in enumerate(datos["productos"], start=1):
        p = item["obj"]
        if p.stock == 0 or p.stock_bajo():
            estilo.add("TEXTCOLOR", (0, i), (-1, i), colors.HexColor("#C0392B"))
    t.setStyle(estilo)
    elementos.append(t)

    doc.build(elementos)
    return _pdf_response(buffer, "reporte_inventario.pdf")


# ============================================================
# 3. PRODUCTOS MÁS VENDIDOS
# ============================================================


def exportar_productos_vendidos_excel(datos):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Más vendidos"

    _titulo_excel(ws, "MOTOPART — Reporte de Productos más Vendidos")
    sub = f"Periodo: {datos['desde']} a {datos['hasta']}"
    sub += " — Categoría: " + (datos["categoria_nombre"] or "Todas")
    ws["A2"] = sub

    fila_tabla = 4
    encabezados = [
        "Puesto",
        "Producto",
        "SKU",
        "Unidades vendidas",
        "Pedidos",
        "Ingresos",
        "Stock actual",
    ]
    for i, h in enumerate(encabezados, start=1):
        ws.cell(row=fila_tabla, column=i, value=h)
    _encabezado_excel(ws, fila_tabla, len(encabezados))

    fila = fila_tabla + 1
    for r in datos["ranking"]:
        ws.cell(row=fila, column=1, value=r["puesto"])
        ws.cell(row=fila, column=2, value=r["nombre"])
        ws.cell(row=fila, column=3, value=r["sku"])
        ws.cell(row=fila, column=4, value=r["unidades_vendidas"])
        ws.cell(row=fila, column=5, value=r["num_pedidos"])
        ws.cell(row=fila, column=6, value=float(r["ingresos"]))
        ws.cell(
            row=fila,
            column=7,
            value=r["stock_actual"] if r["stock_actual"] is not None else "—",
        )
        fila += 1

    if not datos["ranking"]:
        ws.cell(row=fila, column=1, value="No hay ventas registradas en este periodo.")
        fila += 1

    # Segunda tabla: productos sin movimiento
    fila += 2
    ws.cell(
        row=fila,
        column=1,
        value=f"Productos sin movimiento en el periodo ({datos['total_sin_movimiento']})",
    ).font = __import__("openpyxl").styles.Font(bold=True)
    fila += 1
    encabezados2 = ["Producto", "SKU", "Categoría", "Stock actual"]
    for i, h in enumerate(encabezados2, start=1):
        ws.cell(row=fila, column=i, value=h)
    _encabezado_excel(ws, fila, len(encabezados2))
    fila += 1
    for p in datos["sin_movimiento"]:
        ws.cell(row=fila, column=1, value=p.nombre)
        ws.cell(row=fila, column=2, value=p.sku)
        ws.cell(row=fila, column=3, value=p.categoria.nombre if p.categoria else "—")
        ws.cell(row=fila, column=4, value=p.stock)
        fila += 1

    _ajustar_anchos(ws, [9, 30, 14, 16, 10, 16, 12])
    return _xlsx_response(wb, "reporte_productos_vendidos.xlsx")


def exportar_productos_vendidos_pdf(datos):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = _estilo_titulo_pdf()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, topMargin=1.5 * cm, bottomMargin=1.5 * cm
    )
    sub = f"Periodo: {datos['desde']} a {datos['hasta']}"
    sub += " — Categoría: " + (datos["categoria_nombre"] or "Todas")

    elementos = [
        Paragraph("MOTOPART — Reporte de Productos más Vendidos", styles["MPTitulo"]),
        Paragraph(sub, styles["MPSubtitulo"]),
        Spacer(1, 14),
        Paragraph("Ranking de ventas", styles["MPSeccion"]),
    ]

    tabla_datos = [["#", "Producto", "SKU", "Unid.", "Pedidos", "Ingresos", "Stock"]]
    for r in datos["ranking"]:
        tabla_datos.append(
            [
                str(r["puesto"]),
                r["nombre"],
                r["sku"],
                str(r["unidades_vendidas"]),
                str(r["num_pedidos"]),
                f"${_fmt_cop(r['ingresos'])}",
                str(r["stock_actual"]) if r["stock_actual"] is not None else "—",
            ]
        )
    if len(tabla_datos) == 1:
        tabla_datos.append(["—", "Sin ventas en este periodo", "—", "—", "—", "—", "—"])

    t = Table(
        tabla_datos,
        colWidths=[1 * cm, 5.2 * cm, 2.2 * cm, 1.6 * cm, 1.8 * cm, 2.4 * cm, 1.8 * cm],
        repeatRows=1,
    )
    t.setStyle(_tabla_estilo_base())
    elementos.append(t)

    elementos.append(
        Paragraph(
            f"Productos sin movimiento en el periodo ({datos['total_sin_movimiento']})",
            styles["MPSeccion"],
        )
    )
    tabla2 = [["Producto", "SKU", "Categoría", "Stock actual"]]
    for p in datos["sin_movimiento"]:
        tabla2.append(
            [p.nombre, p.sku, p.categoria.nombre if p.categoria else "—", str(p.stock)]
        )
    if len(tabla2) == 1:
        tabla2.append(["—", "—", "Todos tuvieron ventas en este periodo", "—"])

    t2 = Table(tabla2, colWidths=[6 * cm, 2.5 * cm, 4 * cm, 2.5 * cm], repeatRows=1)
    t2.setStyle(_tabla_estilo_base())
    elementos.append(t2)

    doc.build(elementos)
    return _pdf_response(buffer, "reporte_productos_vendidos.pdf")
